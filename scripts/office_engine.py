"""Office converters: docx (mammoth), pptx (python-pptx), xlsx (openpyxl)."""
from __future__ import annotations

from pathlib import Path

from md_utils import demote_headings, md_table


# ------------------------------------------------------------------ docx

def convert_docx(path: Path, log=print):
    """Word document -> markdown via mammoth (docx->semantic HTML) + markdownify."""
    import mammoth

    from web_engine import html_to_md

    with open(path, "rb") as f:
        result = mammoth.convert_to_html(f)
    if result.messages:
        log(f"  (docx: {len(result.messages)} conversion note(s))")
    md = demote_headings(html_to_md(result.value), 1)
    return [(None, md)], {"detail": "Word document (docx)"}


# ------------------------------------------------------------------ pptx

def _shape_md(shape) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            out += _shape_md(sub)
    elif getattr(shape, "has_table", False):
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        table = md_table(rows)
        if table:
            out.append(table)
    elif getattr(shape, "has_chart", False):
        out.append("*[chart]*")
    elif shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
        out.append("*[image]*")
    elif getattr(shape, "has_text_frame", False):
        lines = []
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip() or para.text.strip()
            if not text:
                continue
            level = para.level or 0
            lines.append(f"{'  ' * level}- {text}")
        if lines:
            out.append("\n".join(lines))
    return out


def convert_pptx(path: Path, *, start: int = 1, end: int | None = None, log=print):
    """PowerPoint -> one 'Slide N' section per slide (title, body, tables, notes)."""
    from pptx import Presentation

    prs = Presentation(str(path))
    slides = list(prs.slides)
    total = len(slides)
    start = max(1, start)
    end = min(end or total, total)
    if start > end:
        raise RuntimeError(f"slide range {start}-{end} is empty (deck has {total} slides)")

    sections: list[tuple[str, str]] = []
    for i in range(start - 1, end):
        slide = slides[i]
        pnum = i + 1
        title, title_el = "", None
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text.strip()
                title_el = slide.shapes.title._element
        except Exception:
            pass

        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "_element", None) is title_el:
                continue
            parts += _shape_md(shape)

        notes = ""
        try:
            if slide.has_notes_slide:
                note_text = slide.notes_slide.notes_text_frame.text.strip()
                if note_text:
                    notes = "\n".join(f"> {line}" for line in note_text.splitlines())
        except Exception:
            pass

        body_parts = []
        if title:
            body_parts.append(f"### {title}")
        body_parts += [p for p in parts if p.strip()]
        if notes:
            body_parts.append(notes)
        sections.append((f"Slide {pnum}", "\n\n".join(body_parts)))

    return sections, {"detail": f"{len(sections)} slides (pptx)"}


# ------------------------------------------------------------------ xlsx

def _trim_rows(rows: list[list]) -> list[list]:
    while rows and all(c == "" for c in rows[-1]):
        rows.pop()
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    while width and all(r[width - 1] == "" for r in rows):
        rows = [r[:width - 1] for r in rows]
        width -= 1
    return rows


def convert_xlsx(path: Path, log=print):
    """Excel workbook -> one 'Sheet: name' section per worksheet, rows as pipe tables."""
    import warnings

    from openpyxl import load_workbook

    sections: list[tuple[str, str]] = []
    with warnings.catch_warnings():  # openpyxl "Unknown extension" noise on some files
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(str(path), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [["" if v is None else v for v in row]
                    for row in ws.iter_rows(values_only=True)]
            rows = _trim_rows(rows)
            body = md_table(rows) if rows else "(empty sheet)"
            sections.append((f"Sheet: {ws.title}", body))
        wb.close()
    if not sections:
        sections = [(None, "(no worksheets)")]
    return sections, {"detail": f"{len(sections)} worksheets (xlsx)"}
